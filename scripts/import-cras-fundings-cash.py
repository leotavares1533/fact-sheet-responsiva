from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.util
import json
import math
import re
import unicodedata
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
ENGINE_SPEC = importlib.util.spec_from_file_location("cra_consolidated", ROOT / "scripts/import-cras-carteira-consolidated.py")
ENGINE = importlib.util.module_from_spec(ENGINE_SPEC)
ENGINE_SPEC.loader.exec_module(ENGINE)


def norm(value):
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFD", str(value or "")).encode("ascii", "ignore").decode().lower())


def day(value):
    return str(value)[:10] if re.match(r"^\d{4}-\d{2}-\d{2}", str(value or "")) else ""


def amount(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(value):
        return float(value)
    return None


def prepare_source(source, work):
    import posixpath
    import warnings
    import xml.etree.ElementTree as ET
    import zipfile
    from datetime import time
    import openpyxl

    work.mkdir(parents=True, exist_ok=True)
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
    def convert(value):
        if isinstance(value, (date, datetime, time)):
            return value.isoformat()
        return str(value) if isinstance(value, timedelta) else value
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        values = {sheet.title: [[convert(v) for v in row] for row in sheet.iter_rows(values_only=True)] for sheet in book}
    finally:
        book.close()
    numbers = {}
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(source) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        for sheet in workbook.find("s:sheets", ns):
            rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
            target = targets[rid]
            member = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            cells = {}
            with archive.open(member) as stream:
                for _, cell in ET.iterparse(stream, events=("end",)):
                    if cell.tag == f'{{{ns["s"]}}}c':
                        value = cell.find("s:v", ns)
                        if cell.attrib.get("t", "n") == "n" and value is not None and value.text:
                            cells[cell.attrib["r"]] = float(value.text)
                        cell.clear()
                    elif cell.tag == f'{{{ns["s"]}}}row':
                        cell.clear()
            numbers[sheet.attrib["name"]] = cells
    (work / "source-values.json").write_text(json.dumps(values, ensure_ascii=False), encoding="utf-8")
    (work / "source-raw-numbers.json").write_text(json.dumps(numbers, ensure_ascii=False), encoding="utf-8")


def existing_models():
    models = []
    for folder in sorted((ROOT / "data/cras").glob("cra-carteira-*")):
        for path in sorted(folder.glob("????-??-??.js"), reverse=True):
            snapshot = ENGINE.load_snapshot(path, folder.name, path.stem)
            if snapshot.get("metadata", {}).get("fundingCashOnly"):
                continue
            models.append({"cra": folder.name, "date": path.stem, "metadata": snapshot.get("metadata", {}),
                "cash": snapshot.get("caixa", {}), "funding": snapshot.get("passivo", {}), "craInfo": snapshot.get("cra", {})})
            break
    return models


def br(key):
    return datetime.strptime(key, "%Y-%m-%d").strftime("%d/%m/%Y") if key else ""


def payment_amount(event):
    fields = event["fields"]
    components = fields.get("valorjuros", 0) + fields.get("valoramortizacao", 0)
    return components if components else fields.get("saldoemconta", 0)


def select_payments(entry, issues, decisions):
    grouped = {}
    for event in entry["payments"]:
        grouped.setdefault(event["date"], []).append(event)
    selected = []
    for key, events in sorted(grouped.items()):
        unique = {}
        for event in events:
            unique.setdefault(tuple(sorted(event["fields"].items())), event)
        events = list(unique.values())
        def score(event):
            fields = event["fields"]
            discrepancy = abs(fields.get("saldoemconta", payment_amount(event)) - payment_amount(event))
            quantity = (entry.get("previous") or {}).get("quantidade", 0)
            pu_total = fields.get("pujuros", 0) + fields.get("puamortizacao", 0)
            return (discrepancy > .02, abs(pu_total * quantity - payment_amount(event)) if quantity and pu_total else 0,
                -len(fields))
        events.sort(key=score)
        chosen = copy.deepcopy(events[0])
        conflict = max(payment_amount(e) for e in events) - min(payment_amount(e) for e in events) > .05
        decision_key = f'{entry["cra"]}/{entry["class"]}/{key}'
        decision = decisions.get(decision_key)
        if decision and decision.get("notExecuted"):
            chosen.update({"pending": False, "status": "nao_efetivado", "humanConfirmation": decision})
            selected.append(chosen)
            continue
        if decision:
            chosen["humanConfirmation"] = decision
        chosen["pending"] = conflict and decision is None
        if decision and "amount" in decision:
            chosen = copy.deepcopy(min(events, key=lambda event: abs(payment_amount(event) - decision["amount"])))
            chosen["pending"] = False
            chosen["humanConfirmation"] = decision
        if len(events) > 1:
            issues.append({"cra": entry["cra"], "class": entry["class"], "date": key,
                "type": "quadros_pagamento_repetidos_divergentes" if conflict else "quadros_pagamento_deduplicados",
                "pending": chosen["pending"], "alternatives": events, "selectedAmount": payment_amount(chosen)})
        if not any(abs(chosen["fields"].get(k, 0)) > .00000001 for k in ["valorjuros", "valoramortizacao", "pujuros", "puamortizacao"]):
            chosen["pending"] = not bool(decision and decision.get("reduceFunding"))
            if chosen["pending"]:
                issues.append({"cra": entry["cra"], "class": entry["class"], "date": key,
                    "type": "ajuste_generico_sem_detalhamento" if decision and decision.get("kind") == "outro_ajuste" else "evento_generico_sem_natureza_confirmada",
                    "pending": True, "amount": payment_amount(chosen)})
                if decision:
                    chosen["pendingReason"] = "Outro ajuste confirmado pelo usuario; sem reducao de funding ate detalhamento."
        if chosen["pending"] and any(e.get("dataIso") == key and norm(e.get("tipo")) != "integralizacao"
                for e in (entry.get("previous") or {}).get("eventosAplicados", [])):
            chosen["pending"] = False
            chosen["preservedEarlierValidatedEvent"] = True
            for issue in issues:
                if issue.get("cra") == entry["cra"] and issue.get("class") == entry["class"] and issue.get("date") == key:
                    issue["pending"] = False
                    issue["preservedEarlierValidatedEvent"] = True
        selected.append(chosen)
    return selected


def integral_event(event):
    return {"data": br(event["date"]), "dataIso": event["date"], "tipo": "integralizacao", "quantidade": event["quantity"],
        "valor": event["amount"], "valorEventoReais": event["amount"], "sourceSheet": event["sourceSheet"],
        "sourceRow": event["sourceRow"], "status": event["status"]}


def merge_integrations(entry, cota, cutoff, anchor, issues):
    events = copy.deepcopy(cota.get("eventosAplicados", []))
    known = {e.get("dataIso") or day(e.get("data")): e for e in events if norm(e.get("tipo")) == "integralizacao"}
    additions = []
    for event in entry["integrations"]:
        if event["status"] != "realizado":
            continue
        previous = known.get(event["date"])
        if previous:
            if abs(float(previous.get("quantidade", 0)) - event["quantity"]) > .0001 or abs(float(previous.get("valor", 0)) - event["amount"]) > .02:
                issues.append({"cra": entry["cra"], "class": entry["class"], "date": event["date"],
                    "type": "integralizacao_diverge_documento_preservado", "source": event, "preserved": previous})
            continue
        imported = integral_event(event)
        events.append(imported)
        known[event["date"]] = imported
        additions.append(imported)
        if anchor and event["date"] <= anchor:
            for row in cota.get("historicoPu", []):
                if (row.get("dataIso") or "") >= event["date"]:
                    row["quantidade"] = row.get("quantidade", cota.get("quantidade", 0)) + event["quantity"]
                    row["valorReais"] = row.get("puAtualizado", 0) * row["quantidade"]
            cota["quantidade"] = cota.get("quantidade", 0) + event["quantity"]
    cota["eventosAplicados"] = sorted(events, key=lambda e: e.get("dataIso", ""))
    return additions


def apply_payment(cota, event, entry, issues):
    row = cota["historicoPu"][-1]
    quantity = float(cota.get("quantidade", 0))
    if quantity <= 0:
        raise ValueError(f'Pagamento sem quantidade: {entry["cra"]}/{entry["class"]}/{event["date"]}')
    before = float(row["puAtualizado"])
    paid = payment_amount(event)
    pu_event = paid / quantity
    after = max(0.0, before - pu_event)
    if paid > before * quantity + .02:
        issues.append({"cra": entry["cra"], "class": entry["class"], "date": event["date"],
            "type": "pagamento_maior_saldo_calculado", "pending": True, "payment": paid,
            "balanceBefore": before * quantity, "difference": paid - before * quantity})
        event["excessPaymentReais"] = paid - before * quantity
    row.update({"valorNominal": after, "puAntesEvento": before, "puEvento": pu_event, "puAposEvento": after,
        "puAtualizado": after, "puJuros": 0, "principalAntesEvento": row.get("valorNominal"), "principalAposEvento": after,
        "valorReais": after * quantity, "valorEventoReais": paid, "evento": "Juros / amortizacao" if event["fields"].get("valorjuros") else "Amortizacao / evento",
        "fatorDiAcumulado": 1, "diasUteisPeriodo": 0, "fatorJurosAcumulado": 1,
        "efeitoEvento": "Pagamento deduzido do PU; saldo residual inicia novo periodo"})
    cota.update({"pu": after, "principalResidual": after, "valor": after * quantity})
    item = {"data": br(event["date"]), "dataIso": event["date"], "tipo": "juros_amortizacao", "tipoEvento": row["evento"],
        "valorEventoReais": paid, "valorJuros": event["fields"].get("valorjuros"),
        "valorAmortizacao": event["fields"].get("valoramortizacao"), "puAntesEvento": before,
        "puEvento": pu_event, "puAposEvento": after, "sourceSheet": event["sourceSheet"], "sourceCells": event["sourceCells"]}
    if event.get("excessPaymentReais"):
        item["differencePendingReconciliationReais"] = event["excessPaymentReais"]
    cota.setdefault("eventosAplicados", []).append(item)


def imported_history(entry, cota, events, cutoff):
    first = entry["sourceTerms"]["firstIntegration"]
    integrations = [e for e in entry["integrations"] if e["status"] == "realizado"]
    by_date = {row["date"]: row for row in entry["historyObserved"] if first <= row["date"] <= cutoff}
    previous, history, period = None, [], 0
    for key, observed in sorted(by_date.items()):
        quantity = sum(e["quantity"] for e in integrations if e["date"] <= key)
        pu, principal = observed["pu"], observed["principalPu"]
        if pu is None or principal is None:
            raise ValueError(f'PU sem valor numerico: {entry["cra"]}/{entry["class"]}/{key}')
        today = [e for e in events if e["date"] == key and e["status"] == "realizado" and not e["pending"]]
        if today:
            period = 0
        elif previous:
            period += 1
        paid = sum(payment_amount(e) for e in today)
        history.append({"data": br(key), "dataIso": key, "diaUtil": True, "diasUteis": len(history), "diasUteisPeriodo": period,
            "quantidade": quantity, "valorNominal": principal, "puAtualizado": pu, "puJuros": max(0, pu - principal),
            "fator": pu / principal if principal else 0,
            "puAntesEvento": pu + paid / quantity if quantity else pu, "puEvento": paid / quantity if quantity else 0,
            "puAposEvento": pu, "principalAposEvento": principal, "valorReais": pu * quantity, "valorEventoReais": paid,
            "evento": "Juros / amortizacao" if today else "Integralizacao" if any(e["date"] == key for e in integrations) else "",
            "taxaDiStatus": "importada_pendente_validacao", "sourceRow": observed["sourceRow"]})
        previous = observed
    if not history or history[-1]["dataIso"] != cutoff:
        raise ValueError(f'Historico nao termina na data-base: {entry["cra"]}/{entry["class"]}')
    cota.update({"historicoPu": history, "pu": history[-1]["puAtualizado"], "principalResidual": history[-1]["valorNominal"],
        "quantidade": history[-1]["quantidade"], "valor": history[-1]["valorReais"]})
    for event in events:
        if event["status"] != "realizado" or event["pending"]:
            continue
        row = next(r for r in history if r["dataIso"] == event["date"])
        cota["eventosAplicados"].append({"data": br(event["date"]), "dataIso": event["date"], "tipo": "juros_amortizacao",
            "valorEventoReais": payment_amount(event), "valorJuros": event["fields"].get("valorjuros"),
            "valorAmortizacao": event["fields"].get("valoramortizacao"), "puEvento": row["puEvento"],
            "puAposEvento": row["puAtualizado"], "sourceSheet": event["sourceSheet"], "sourceCells": event["sourceCells"]})


def build_cota(entry, model, cutoff, issues, decisions):
    validated = entry["existingValidatedSeries"]
    old = copy.deepcopy(entry.get("previous") or {})
    terms = entry["sourceTerms"]
    cota = old if validated else {"classe": entry["class"], "tipo": "sr" if entry["class"] == "SR1" else "mez",
        "label": {"SR1": "Senior", "MEZA": "Mezanino A", "MEZB": "Mezanino B"}[entry["class"]], "ifCodigo": old.get("ifCodigo", ""),
        "ehFunding": True, "quantidade": 0, "valorNominalInicial": 1000, "metodo": terms["method"],
        "taxaAa": terms["annualRate"] or 0, "spreadAa": terms["spreadAa"] or 0,
        "indexador": "DI" if terms["spreadAa"] is not None else "", "percentualIndexador": 1,
        "dataInicioIso": terms["firstIntegration"], "dataInicio": br(terms["firstIntegration"]),
        "dataVencimentoIso": terms["maturity"] if "vencimento_anterior_integralizacao" not in entry["checks"] else "",
        "dataVencimento": br(terms["maturity"]) if "vencimento_anterior_integralizacao" not in entry["checks"] else "",
        "revolvencia": None, "eventosAplicados": [], "historicoPu": []}
    if validated and "di_spread" in cota.get("metodo", ""):
        spread = ENGINE.extract_spread_from_text(cota.get("taxaTexto", ""))
        if spread is not None:
            cota["spreadAa"] = spread
    cota["taxaTexto"] = rate_text(cota)
    payments = select_payments(entry, issues, decisions)
    anchor = model.get("date", "") if validated and cota.get("quantidade", 0) else ""
    additions = merge_integrations(entry, cota, cutoff, anchor, issues)
    if not validated:
        imported_history(entry, cota, payments, cutoff)
    else:
        if not anchor:
            first = terms["firstIntegration"]
            initial = next(e for e in additions if e["dataIso"] == first)
            cota["quantidade"] = initial["quantidade"]
            cota["historicoPu"] = [{"data": br(first), "dataIso": first, "valorNominal": 1000,
                "puAtualizado": 1000, "valorReais": 1000 * initial["quantidade"], "quantidade": initial["quantidade"],
                "diasUteis": 0, "diasUteisPeriodo": 0, "fatorDiAcumulado": 1, "produtorioFatorDi": 1,
                "evento": "Integralizacao", "valorEventoReais": initial["valor"]}]
            anchor = first
        timeline = sorted({e["dataIso"] for e in additions if anchor < e["dataIso"] <= cutoff}
            | {e["date"] for e in payments if anchor < e["date"] <= cutoff and e["status"] == "realizado" and not e["pending"]})
        context = {"metadata": copy.deepcopy(model.get("metadata", {})), "performanceCotas": []}
        for key in timeline:
            cota = ENGINE.project_cota_to_date(context, cota, key)
            today = [e for e in additions if e["dataIso"] == key]
            if today:
                cota["quantidade"] += sum(e["quantidade"] for e in today)
                row = cota["historicoPu"][-1]
                row.update({"quantidade": cota["quantidade"], "valorReais": row["puAtualizado"] * cota["quantidade"],
                    "evento": "Integralizacao", "valorEventoReais": sum(e["valor"] for e in today)})
            for event in payments:
                if event["date"] == key and event["status"] == "realizado" and not event["pending"]:
                    apply_payment(cota, event, entry, issues)
        cota = ENGINE.project_cota_to_date(context, cota, cutoff)
        for event in payments:
            if event["status"] == "realizado" and event["date"] <= anchor and not any(e.get("dataIso") == event["date"] and norm(e.get("tipo")) != "integralizacao" for e in old.get("eventosAplicados", [])):
                event["pending"] = True
                event["pendingReason"] = "Evento anterior a memoria validada; conciliar comunicado antes de substituir historico."
                issues.append({"cra": entry["cra"], "class": entry["class"], "date": event["date"], "type": "evento_anterior_base_validada_pendente", "pending": True})
    source_balance = (entry.get("latestObserved") or {}).get("balance")
    observed = entry.get("latestObserved") or {}
    if source_balance is not None and observed.get("pu") is not None and abs(source_balance - observed["pu"] * cota["quantidade"]) > .02:
        issues.append({"cra": entry["cra"], "class": entry["class"], "date": cutoff, "type": "saldo_planilha_nao_e_pu_vezes_quantidade",
            "sourceBalance": source_balance, "puTimesQuantity": observed["pu"] * cota["quantidade"]})
    cota["documentValidation"] = {"status": "validado_cadastro_anterior" if validated else "pendente_documentos",
        "fields": [] if validated else ["codigo_if", "taxa", "datas", "prazo", "revolvencia", "eventos"],
        "sourceTerms": terms, "checks": entry["checks"]}
    cota["dataPrimeiraIntegralizacaoIso"] = terms["firstIntegration"]
    if cota.get("dataVencimentoIso") and cota["dataVencimentoIso"] < cutoff and cota["valor"] > .02:
        issues.append({"cra": entry["cra"], "class": entry["class"], "date": cutoff,
            "type": "saldo_apos_vencimento_cadastrado", "pending": True, "maturity": cota["dataVencimentoIso"], "balance": cota["valor"]})
    cota["importacaoFundingCaixa"] = {"sourceSheet": entry["sheet"], "sourcePu": observed.get("pu"),
        "sourceBalance": source_balance, "calculation": "taxas_contratuais_preservadas" if validated else "pu_planilha_pendente_validacao"}
    cota["forecastBlockedReason"] = "Agenda futura e documentos pendentes de conciliacao nesta importacao."
    cota["previsaoPu"] = []
    cota["eventosPendentes"] = [e for e in payments if e["pending"]]
    cota["eventosPrevistos"] = [e for e in payments if e["status"] == "previsto"]
    cota["eventosAplicados"].sort(key=lambda e: e.get("dataIso", ""))
    subscriptions = [e for e in cota["eventosAplicados"] if norm(e.get("tipo")) == "integralizacao"]
    for row in cota["historicoPu"]:
        quantity = sum(e.get("quantidade", 0) for e in subscriptions if e.get("dataIso", "") <= row["dataIso"])
        row["quantidade"] = quantity
        row["valorReais"] = row["puAtualizado"] * quantity
    cota["dataHistoricaDisponivel"] = True
    cota["acumulacaoFinal"] = {"periodoInicio": cota.get("dataInicio"), "periodoFim": br(cutoff), "puFinal": cota["pu"],
        "diasAcumulacao": cota["historicoPu"][-1].get("diasUteis", 0)}
    return cota


def apply_report(report, models, work, decisions):
    cutoff = report["positionDate"]
    registry_path = ROOT / "data/fundings/cras-carteira/documented-terms.json"
    if registry_path.exists():
        decisions = copy.deepcopy(decisions)
        registry = json.loads(registry_path.read_text(encoding="utf-8"))
        for decision in registry.get("eventDecisions", []):
            if decision.get("status") == "nao_efetivado":
                for scope in decision["scope"]:
                    decisions[f'{scope}/{decision["date"]}'] = {"notExecuted": True, "confirmedBy": decision["confirmedBy"], "confirmationDate": decision["confirmationDate"]}
    by_cra = {row["cra"]: row for row in models}
    issues = report["issues"]
    snapshots = {}
    for cra in sorted(set(report["cash"]) | {row["cra"] for row in report["series"]}, key=lambda s: int(s.rsplit("-", 1)[-1])):
        model = by_cra.get(cra, {})
        entries = [e for e in report["series"] if e["cra"] == cra]
        cotas = [build_cota(e, model, cutoff, issues, decisions) for e in entries]
        for old in model.get("funding", {}).get("cotas", []):
            if old.get("ehFunding") and not any(c["classe"] == old["classe"] for c in cotas):
                preserved = copy.deepcopy(old)
                preserved["forecastBlockedReason"] = "Serie ausente no arquivo novo; cadastro anterior preservado."
                cotas.append(preserved)
        cash = report["cashCurrent"].get(cra)
        if not cash:
            issues.append({"cra": cra, "type": "caixa_nao_informado", "pending": True})
        if not entries:
            issues.append({"cra": cra, "type": "funding_nao_informado", "pending": True})
        metadata = {"schemaVersion": 2, "model": "cra-carteira", "craId": cra, "dateKey": cutoff, "reportDate": br(cutoff),
            "revisionId": f'{cutoff.replace("-", "")}-funding-caixa-{report["sourceHash"][:10]}',
            "importedAt": datetime.now(timezone(timedelta(hours=-3))).isoformat(), "fundingCashOnly": True,
            "crasCarteiraBatch": True, "businessCalendar": copy.deepcopy(model.get("metadata", {}).get("businessCalendar") or models[0]["metadata"]["businessCalendar"]),
            "previousPortfolioDateKey": model.get("date"), "fundingCashImport": {"sourceFile": Path(report["source"]).name,
                "sourceHash": report["sourceHash"], "sheetCash": "Saldo Caixa", "documentValidationRequired": True}}
        snapshot = {"metadata": metadata, "cra": {"id": cra, "name": f'CRA Carteira {cra.rsplit("-", 1)[-1]}',
                "dateKey": cutoff, "dataBase": br(cutoff)},
            "ativo": {"total": None, "caixa": cash["total"] if cash else None},
            "caixa": {"accounts": {"cc": cash.get("cc", 0), "contaAplicacao": cash.get("aplicacao", 0),
                "fundoDespesas": cash.get("despesa", 0)} if cash else {}, "total": cash["total"] if cash else None,
                "totalCalculado": cash["calculated"] if cash else None, "fonte": "Saldo Caixa", "arquivoOrigem": Path(report["source"]).name,
                "observacao": "Total informado preservado. Fundo de despesas deduzido uma unica vez. Divergencias sinalizadas."},
            "passivo": {"fundingTotal": sum(c.get("valor", 0) for c in cotas) if cotas else None,
                "subordinadaTotal": None, "subordinadaPuResidual": None, "cotas": cotas},
            "sources": [{"name": "Eventos de funding e caixa", "date": br(cutoff), "file": Path(report["source"]).name}],
            "caixaHistorico": {key: {"total": row["total"], "totalCalculado": row["calculated"],
                "accounts": {"cc": row.get("cc", 0), "contaAplicacao": row.get("aplicacao", 0), "fundoDespesas": row.get("despesa", 0)}}
                for key, row in report["cash"].get(cra, {}).items()},
            "fundingCashIssues": []}
        snapshots[cra] = snapshot
    for cra, snapshot in snapshots.items():
        snapshot["fundingCashIssues"] = [e for e in issues if e.get("cra") == cra
            and (e.get("type") != "caixa_nao_concilia" or e.get("date") == cutoff)]
    overview_rows = [{"operacao": s["cra"]["name"], "craId": cra, "serie": c.get("label") or c["classe"], "classe": c["classe"],
        "ifCodigo": c.get("ifCodigo"), "quantidadeIntegralizada": c.get("quantidade"), "puAtual": c.get("pu"), "valorAtual": c.get("valor"),
        "taxa": rate_text(c), "dataVencimento": c.get("dataVencimento"), "dataVencimentoIso": c.get("dataVencimentoIso"),
        "status": "Pendente documentos" if c.get("documentValidation", {}).get("status") == "pendente_documentos" else "Integralizada",
        "documentValidation": c.get("documentValidation")} for cra, s in snapshots.items() for c in s["passivo"]["cotas"]]
    overview = {"metadata": {**snapshots[next(iter(snapshots))]["metadata"], "portfolioOverview": True,
            "craId": "cras-carteira-overview"}, "cra": {"id": "cras-carteira-overview", "name": "CRAs Carteira - Visao geral"},
        "ativo": {"total": None, "caixa": sum(r["total"] for r in report["cashCurrent"].values())},
        "caixa": {"total": sum(r["total"] for r in report["cashCurrent"].values())},
        "passivo": {"fundingTotal": sum(s["passivo"]["fundingTotal"] or 0 for s in snapshots.values()), "cotas": []},
        "portfolioOverview": {"rows": overview_rows}, "fundingCashVehicles": [{"craId": cra, "name": s["cra"]["name"],
            "funding": s["passivo"]["fundingTotal"], "cash": s["caixa"]["total"]} for cra, s in snapshots.items()],
        "fundingCashIssues": [{"cra": e["cra"], "class": e.get("class"), "date": e.get("date"), "type": e["type"]}
            for e in issues if e.get("pending") or e.get("type") == "caixa_nao_concilia" and e.get("date") == cutoff],
        "sources": snapshots[next(iter(snapshots))]["sources"]}
    registry_path = ROOT / "data/fundings/cras-carteira/documented-terms.json"
    if registry_path.exists():
        review_spec = importlib.util.spec_from_file_location("cra_carteira_document_review", ROOT / "scripts/reconcile-cras-carteira-documents.py")
        review = importlib.util.module_from_spec(review_spec)
        review_spec.loader.exec_module(review)
        registry = json.loads(registry_path.read_text(encoding="utf-8"))
        snapshots, document_changes = review.apply_document_review(snapshots, report, registry)
        overview = review.build_overview(overview, snapshots, registry["reviewDate"])
        report["documentReview"] = {"date": registry["reviewDate"], "registry": "data/fundings/cras-carteira/documented-terms.json",
            "cancelledPaymentDate": "2026-10-07", "seriesReviewed": sum(c["documented"] for c in document_changes)}
    report["summary"] = {"positionDate": cutoff, "series": len(report["series"]), "cashVehicles": len(report["cashCurrent"]),
        "vehicles": len(snapshots), "newSeries": [{"cra": e["cra"], "class": e["class"]} for e in report["series"] if not e["existingValidatedSeries"]],
        "eventsApplied": sum(len(c.get("eventosAplicados", [])) for s in snapshots.values() for c in s["passivo"]["cotas"]),
        "pendingPayments": sum(len(c.get("eventosPendentes", [])) for s in snapshots.values() for c in s["passivo"]["cotas"]),
        "cashTotal": overview["caixa"]["total"], "fundingTotal": overview["passivo"]["fundingTotal"]}
    report["currentPositions"] = overview["fundingCashVehicles"]
    data_folder = ROOT / "data/fundings/cras-carteira"
    data_folder.mkdir(parents=True, exist_ok=True)
    (data_folder / f"{cutoff}.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    for cra, snapshot in {**snapshots, "cras-carteira-overview": overview}.items():
        ENGINE.write_snapshot(ROOT / f"data/cras/{cra}/{cutoff}.js", cra, cutoff, snapshot)
    ENGINE.update_manifest(cutoff, snapshots, overview)
    # Funding-only dates must never be labelled as portfolio imports.
    manifest_path = ROOT / "data/cra-manifest.js"
    manifest = ENGINE.load_json_assignment(manifest_path)
    for item in manifest:
        if item["craId"] in snapshots or item["craId"] == "cras-carteira-overview":
            for row in item["dates"]:
                if row["dateKey"] == cutoff:
                    row.update({"fundingCashOnly": True, "totalAtivo": None, "subordinada": None,
                        "assetImported": False, "portfolioAssetImport": False})
    ENGINE.write_manifest(manifest_path, manifest)
    write_report(report)
    print(json.dumps(report["summary"], indent=2, ensure_ascii=True))


def rate_text(cota):
    if cota.get("taxaTexto"):
        return cota["taxaTexto"]
    spread = cota.get("spreadAa", 0)
    return f'CDI + {spread * 100:.2f}% a.a.' if "di_spread" in cota.get("metodo", "") else f'{cota.get("taxaAa", 0) * 100:.2f}% a.a.'


def write_report(report):
    folder = ROOT / "output/relatorios"
    folder.mkdir(parents=True, exist_ok=True)
    def money(value):
        return "Nao informado" if value is None else "R$ " + f"{value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    def cra_name(value):
        return "CRA " + value.rsplit("-", 1)[-1]
    lines = [f'# Atualizacao dos fundings e caixa - {br(report["positionDate"])}', "", f'Fonte: {Path(report["source"]).name}',
        f'SHA-256: {report["sourceHash"]}', "", "## Escopo", "",
        f'{report["summary"]["series"]} series e {report["summary"]["cashVehicles"]} caixas importados. Historico anterior preservado em seus snapshots.',
        "Sem importacao da carteira de ativos: PL, PDD e subordinada nao recalculados. CRAs 42 e 65 nao alterados.",
        "Taxas e prazos documentados preservados. Novas series usam o PU observado da planilha, sem validacao contratual. Projecoes futuras desta importacao bloqueadas ate conciliacao da agenda.",
        "", "## Documentos necessarios", "", "Enviar TS, aditamentos e comunicados para confirmar IF, taxa, datas, vencimento, revolvencia e eventos.", "",
        "| Emissao | Serie | Remuneracao informada, nao validada | Vencimento informado |", "|---|---|---|---|"]
    for e in report["series"]:
        if not e["existingValidatedSeries"]:
            t = e["sourceTerms"]
            rate = f'CDI + {t["spreadAa"] * 100:.2f}% a.a.' if t["spreadAa"] is not None else f'{t["annualRate"] * 100:.2f}% a.a.'
            lines.append(f'| {e["cra"].rsplit("-", 1)[-1]} | {e["class"]} | {rate} | {br(t["maturity"]) or "Nao informado"} |')
    lines += ["", "## Posicao atual", "", "Os saldos abaixo incluem valores provisorios e pagamentos ainda pendentes de conciliacao.", "",
        "| Veiculo | Funding SR/MEZ | Caixa informado |", "|---|---:|---:|"]
    for row in report["currentPositions"]:
        lines.append(f'| {row["name"]} | {money(row["funding"])} | {money(row["cash"])} |')
    lines += ["", "## Conferencias pendentes", "", "Valores de caixa informados preservados, mesmo quando o total foi limitado a zero na planilha.", "",
        "| Veiculo / serie | Data | Ponto a conferir |", "|---|---|---|"]
    for issue in report["issues"]:
        if issue["type"] == "caixa_nao_concilia" and issue["date"] != report["positionDate"]:
            continue
        if issue["type"] == "quadros_pagamento_deduplicados":
            continue
        texts = {
            "caixa_nao_concilia": f'Total informado {money(issue.get("reported"))}; contas menos despesas {money(issue.get("calculated"))}.',
            "caixa_nao_informado": "Caixa ausente no arquivo; nao presumido como zero.",
            "funding_nao_informado": "Somente caixa informado; falta a memoria do funding.",
            "pagamento_maior_saldo_calculado": f'Pagamento de {money(issue.get("payment"))} excede o saldo calculado em {money(issue.get("difference"))}. Pagamento registrado integralmente; saldo do funding limitado a zero, diferenca separada para conciliacao. Nao tratada como receita ou despesa.',
            "evento_generico_sem_natureza_confirmada": f'Evento generico de {money(issue.get("amount"))}: confirmar natureza e efeito no PU.',
            "ajuste_generico_sem_detalhamento": f'Outro ajuste de {money(issue.get("amount"))}, conforme confirmacao de Leonardo. Sem reducao de funding como pagamento; aguarda detalhamento.',
            "evento_anterior_base_validada_pendente": "Quadro de pagamento antigo ausente na memoria validada: confirmar comunicado antes de substituir o historico.",
            "quadros_pagamento_repetidos_divergentes": "Dois quadros de pagamento divergentes na mesma data. Nenhum pagamento adicional presumido.",
            "integralizacao_diverge_documento_preservado": "Quantidade/valor da integralizacao diverge do cadastro documentado; preservado o documento, sem duplicar a entrada.",
            "saldo_planilha_nao_e_pu_vezes_quantidade": f'Coluna de saldo total da planilha {money(issue.get("sourceBalance"))} difere de PU x quantidade da serie {money(issue.get("puTimesQuantity"))}. A coluna agregada nao foi usada como saldo individual.',
            "saldo_apos_vencimento_cadastrado": f'Havera saldo de {money(issue.get("balance"))} depois do vencimento cadastrado ({br(issue.get("maturity", ""))}); conferir prazo e pagamentos.'}
        text = texts.get(issue["type"], issue["type"])
        if issue.get("preservedEarlierValidatedEvent"):
            text += " Evento ja validado anteriormente foi preservado."
        lines.append(f'| {cra_name(issue["cra"])} {issue.get("class", "")} | {br(issue.get("date", "")) or "-"} | {text} |')
    lines += ["", "## Divergencias de cadastro", ""]
    for e in report["series"]:
        if e["checks"]:
            lines.append(f'- {e["cra"]} {e["class"]}: {", ".join(e["checks"])}.')
    lines += ["", "O detalhamento auditavel, inclusive celulas de origem e historico completo de caixa, esta em data/fundings/cras-carteira/" + report["positionDate"] + ".json.", ""]
    (folder / f'Atualizacao_Fundings_Caixa_CRAs_Carteira_{report["positionDate"]}.md').write_text("\n".join(lines), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description="Importa eventos e caixa dos CRAs Carteira sem substituir taxas contratuais validadas.")
    parser.add_argument("--prepared-source", default="tmp/fundings-carteiras-import")
    parser.add_argument("--source", required=True)
    parser.add_argument("--apply", action="store_true", help="Atualiza snapshots, cadastro e caixa; sem publicar ou comitar.")
    parser.add_argument("--decisions", help="JSON de confirmacoes humanas para eventos ambiguos.")
    args = parser.parse_args()
    work = Path(args.prepared_source)
    source_hash = hashlib.sha256(Path(args.source).read_bytes()).hexdigest()
    hash_path = work / "source-hash.txt"
    if not (work / "source-values.json").exists() or not hash_path.exists() or hash_path.read_text().strip() != source_hash:
        prepare_source(Path(args.source), work)
        hash_path.write_text(source_hash, encoding="ascii")
    values = json.loads((work / "source-values.json").read_text(encoding="utf-8"))
    numbers = json.loads((work / "source-raw-numbers.json").read_text(encoding="utf-8"))
    models = existing_models()
    existing = {(row["cra"], cota["classe"]): cota for row in models for cota in row.get("funding", {}).get("cotas", []) if cota.get("ehFunding")}

    def raw(sheet, row, col):
        return numbers.get(sheet, {}).get(f"{get_column_letter(col + 1)}{row + 1}", amount(values[sheet][row][col]) if col < len(values[sheet][row]) else None)

    cash_rows = values["Saldo Caixa"]
    # Cash values have accidental date formats; only the dated header is a calendar.
    def cash_header(row):
        if len(row) <= 6 or norm(row[5]):
            return {}
        return {j: day(value) for j, value in enumerate(row) if j >= 6 and day(value)
            and "2025-01-01" <= day(value) <= datetime.now().date().isoformat()}

    cash_dates = [value for row in cash_rows for value in cash_header(row).values()]
    cutoff = max(cash_dates)
    cash, cash_current, date_columns, current_cra = {}, {}, {}, ""
    issues, series = [], []
    for i, row in enumerate(cash_rows):
        dates = cash_header(row)
        if dates:
            date_columns = dates
        label = str(row[1] or "") if len(row) > 1 else ""
        match = re.fullmatch(r"(\d+)\s*[ªaºo]?", label.strip())
        if match:
            current_cra = f"cra-carteira-{match.group(1)}"
        if not current_cra or len(row) <= 5:
            continue
        field = norm(row[5])
        name = "cc" if "contacorrente" in field else "aplicacao" if "aplicacao" in field else "despesa" if "fundodedespesa" in field else "total" if field == "total" else ""
        if not name:
            continue
        for j, date_key in date_columns.items():
            value = raw("Saldo Caixa", i, j)
            if value is None:
                continue
            record = cash.setdefault(current_cra, {}).setdefault(date_key, {"date": date_key, "sourceCells": {}})
            record[name] = value
            record["sourceCells"][name] = f"{get_column_letter(j + 1)}{i + 1}"
    for cra, records in cash.items():
        for date_key, record in records.items():
            if "total" not in record:
                record["total"] = record.get("cc", 0) + record.get("aplicacao", 0) - record.get("despesa", 0)
            calculated = record.get("cc", 0) + record.get("aplicacao", 0) - record.get("despesa", 0)
            record["calculated"] = calculated
            if abs(record["total"] - calculated) > .011:
                issues.append({"cra": cra, "type": "caixa_nao_concilia", "date": date_key,
                    "reported": record["total"], "calculated": calculated, "difference": record["total"] - calculated})
        cash_current[cra] = records[max(records)]

    for sheet, rows in values.items():
        if sheet == "Saldo Caixa":
            continue
        match = re.match(r"(\d+).*?(SR|MEZ\s*[AB])", sheet, re.I)
        if not match:
            raise ValueError(f"Aba nao identificada: {sheet}")
        cra = f"cra-carteira-{match.group(1)}"
        classe = "SR1" if match.group(2).upper() == "SR" else re.sub(r"\s", "", match.group(2)).upper()
        old = existing.get((cra, classe))
        known = bool(old and old.get("ifCodigo") and old.get("metodo") != "pendente_integralizacao")
        date_rows = [(i, row) for i, row in enumerate(rows) if len(row) > 1 and day(row[1])]
        first_daily = date_rows[0][0]
        integrations, pending_integrations = [], []
        for i, row in enumerate(rows[:first_daily]):
            if len(row) < 5 or not day(row[2]):
                continue
            paid, quantity = raw(sheet, i, 3), raw(sheet, i, 4)
            if paid is None or quantity is None:
                if paid is not None or quantity is not None or norm(row[1]) != "vencimento":
                    pending_integrations.append({"date": day(row[2]), "row": i + 1, "amount": paid, "quantity": quantity})
                continue
            if paid <= 0 or quantity <= 0:
                continue
            if abs(quantity - round(quantity)) < .0001:
                quantity = float(round(quantity))
            integrations.append({"date": day(row[2]), "type": "integralizacao", "amount": paid, "quantity": quantity,
                "sourceRow": i + 1, "sourceSheet": sheet, "status": "realizado" if day(row[2]) <= cutoff else "previsto"})
        payments = []
        for i, row in enumerate(rows):
            for j, label in enumerate(row[:-1]):
                if norm(label) not in {"amortizacao", "pagamentodejuros", "pagamentojuros", "pagamento"} or not day(row[j + 1]):
                    continue
                date_key = day(row[j + 1])
                fields, locations = {}, {}
                for k in range(i + 1, min(i + 12, len(rows))):
                    if j + 1 >= len(rows[k]):
                        continue
                    key = norm(rows[k][j])
                    if key == "amortizacao" and day(rows[k][j + 1]):
                        break
                    value = raw(sheet, k, j + 1)
                    if value is not None and key in {"saldoemconta", "pujuros", "valorjuros", "valoramortizacao", "puamortizacao", "saldodevedor", "pueventogenerico"}:
                        fields[key], locations[key] = value, f"{get_column_letter(j + 2)}{k + 1}"
                if not any(abs(value) > .00000001 for key, value in fields.items() if key != "saldodevedor"):
                    continue
                payments.append({"date": date_key, "type": "pagamento", "sourceRow": i + 1, "sourceColumn": j + 1,
                    "sourceSheet": sheet, "fields": fields, "sourceCells": locations,
                    "status": "realizado" if date_key <= cutoff else "previsto"})
        historical = [{"date": day(row[1]), "sourceRow": i + 1, "principalPu": raw(sheet, i, 2),
            "pu": raw(sheet, i, 4), "balance": raw(sheet, i, 7)} for i, row in date_rows if day(row[1]) <= cutoff]
        maturity = next((day(row[2]) for row in rows[:first_daily] if len(row) > 2 and norm(row[1]) == "vencimento"), "")
        header_index = next(i for i, row in enumerate(rows[:first_daily]) if len(row) > 4 and norm(row[1]) == "data" and norm(row[4]) == "puatualizado")
        fixed = "juros" in norm(rows[header_index][3])
        source_rate = raw(sheet, header_index + 1, 3)
        latest = historical[-1] if historical else None
        source_terms = {"method": "prefixado_252_ts" if fixed else "di_spread_252", "annualRate": source_rate if fixed else None,
            "spreadAa": source_rate if not fixed else None, "maturity": maturity, "firstIntegration": min((row["date"] for row in integrations), default="")}
        checks = []
        if not known:
            checks.append("validar_documentos_taxa_prazo_revolvencia")
        if not maturity:
            checks.append("vencimento_ausente")
        if maturity and integrations and maturity < min(row["date"] for row in integrations):
            checks.append("vencimento_anterior_integralizacao")
        if known and old.get("metodo") != source_terms["method"]:
            checks.append("taxa_planilha_diverge_cadastro_validado")
        if known and maturity and maturity != old.get("dataVencimentoIso"):
            checks.append("vencimento_planilha_diverge_cadastro_validado")
        for item in pending_integrations:
            checks.append("integralizacao_sem_valor_ou_quantidade")
        entry = {"cra": cra, "class": classe, "sheet": sheet, "existingValidatedSeries": known, "previous": old,
            "sourceTerms": source_terms, "integrations": integrations, "payments": payments, "historyObserved": historical,
            "pendingIntegrations": pending_integrations, "checks": sorted(set(checks)), "latestObserved": latest}
        series.append(entry)
    report = {"source": str(Path(args.source).resolve()), "sourceHash": hashlib.sha256(Path(args.source).read_bytes()).hexdigest(),
        "positionDate": cutoff, "series": series, "cash": cash, "cashCurrent": cash_current, "issues": issues}
    (work / "preflight.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    if args.apply:
        protection = work / "protected-before.json"
        if not protection.exists():
            paths = [p for p in (ROOT / "data/cras").glob("*/*.js") if p.stem != cutoff or not (p.parent.name.startswith("cra-carteira-") or p.parent.name == "cras-carteira-overview")]
            protection.write_text(json.dumps({str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}), encoding="utf-8")
            (work / "manifest-before.js").write_bytes((ROOT / "data/cra-manifest.js").read_bytes())
        decisions_path = Path(args.decisions) if args.decisions else ROOT / "data/fundings/cras-carteira/confirmacoes-2026-10-29.json"
        decisions = json.loads(decisions_path.read_text(encoding="utf-8")) if decisions_path.exists() else {}
        apply_report(report, models, work, decisions)
        return
    print(json.dumps({"positionDate": cutoff, "seriesCount": len(series), "cashVehicleCount": len(cash),
        "cashIssues": [row for row in issues if row["date"] == cutoff], "historicalCashIssueCount": len(issues),
        "series": [{"cra": row["cra"], "class": row["class"], "new": not row["existingValidatedSeries"],
            "integrations": len(row["integrations"]), "payments": len(row["payments"]), "checks": row["checks"],
            "latest": row["latestObserved"]} for row in series]}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
